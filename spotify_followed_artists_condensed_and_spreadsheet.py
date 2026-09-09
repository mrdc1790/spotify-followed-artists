import re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule

# =========================================================
# INPUT FILE
# =========================================================

INPUT_FILE = "spotify_followed_artists_20260729_221839.txt"

# =========================================================
# READ FILE
# =========================================================

with open(INPUT_FILE, "r", encoding="utf-8") as f:
    content = f.read()

# =========================================================
# PARSE ARTISTS
# =========================================================

pattern = re.compile(
    r"(\d+)\.\s(.+?)\n"
    r"\s+Followers:\s([\d,]+)\n"
    r"\s+Popularity:\s(\d+)\n"
    r"\s+Genres:\s(.*?)\n"
    r"\s+URL:\s(https://open\.spotify\.com/artist/[^\s]+)",
    re.DOTALL
)

matches = pattern.findall(content)

artists = []

for match in matches:

    number, name, followers, popularity, genres, url = match

    artists.append({
        "number": int(number),
        "name": name.strip(),
        "followers": int(followers.replace(",", "")),
        "popularity": int(popularity),
        "genres": genres.strip(),
        "url": url.strip()
    })

print(f"Parsed {len(artists)} artists.")

# =========================================================
# OUTPUT 1: NAMES ONLY TXT
# =========================================================

names_output = Path(INPUT_FILE).stem + "_names_only.txt"

with open(names_output, "w", encoding="utf-8") as f:

    for artist in artists:
        f.write(artist["name"] + "\n")

print(f"Created: {names_output}")

# =========================================================
# OUTPUT 2: EXCEL FILE
# =========================================================

wb = Workbook()
ws = wb.active

ws.title = "Spotify Artists"

headers = [
    "#",
    "Artist",
    "Followers",
    "Popularity",
    "Genres",
    "Spotify URL"
]

ws.append(headers)

# =========================================================
# HEADER STYLING
# =========================================================

header_fill = PatternFill(
    start_color="1DB954",
    end_color="1DB954",
    fill_type="solid"
)

header_font = Font(
    bold=True,
    color="FFFFFF"
)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font

# =========================================================
# DATA
# =========================================================

for artist in artists:

    row = [
        artist["number"],
        artist["name"],
        artist["followers"],
        artist["popularity"],
        artist["genres"],
        artist["url"]
    ]

    ws.append(row)

# =========================================================
# HYPERLINKS
# =========================================================

for row in range(2, ws.max_row + 1):

    cell = ws[f"F{row}"]

    cell.hyperlink = cell.value
    cell.style = "Hyperlink"

# =========================================================
# FILTERS + FREEZE
# =========================================================

ws.auto_filter.ref = ws.dimensions
ws.freeze_panes = "A2"

# =========================================================
# COLUMN WIDTHS
# =========================================================

widths = {
    1: 8,
    2: 35,
    3: 14,
    4: 12,
    5: 50,
    6: 60
}

for col_num, width in widths.items():
    ws.column_dimensions[get_column_letter(col_num)].width = width

# =========================================================
# CONDITIONAL FORMATTING
# =========================================================

popularity_range = f"D2:D{ws.max_row}"

ws.conditional_formatting.add(
    popularity_range,
    ColorScaleRule(
        start_type='min',
        start_color='FFF5F5',
        mid_type='percentile',
        mid_value=50,
        mid_color='FFE599',
        end_type='max',
        end_color='93C47D'
    )
)

# =========================================================
# SAVE
# =========================================================

excel_output = Path(INPUT_FILE).stem + ".xlsx"

wb.save(excel_output)

print(f"Created: {excel_output}")