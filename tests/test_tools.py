"""Offline regression tests: no Spotify login or account changes."""

import importlib.util
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from electronic_artists_edm import genre_report, group_by_genre
from hipHopArtists import get_all_and_filter_followed_artists
from playlist_tools import collect_tracks, create_playlist
from spotify_common import followed_artists
from spotify_followed_artists_list import export_artists
from spotify_followed_artists_condensed_and_spreadsheet import convert_export, parse_export
from unique_artists import extract_names
from unsaveDuplicates import find_duplicates, saved_tracks


class ToolTests(unittest.TestCase):
    def test_cursor_pagination_including_short_page(self):
        client = Mock()
        client.current_user_followed_artists.side_effect = [
            {"artists": {"items": [{"name": "a"}], "cursors": {"after": "next"}}},
            {"artists": {"items": [{"name": "B"}], "cursors": {"after": None}}},
        ]
        self.assertEqual(get_all_and_filter_followed_artists(client, "b"), ["B"])
        self.assertEqual(client.current_user_followed_artists.call_count, 2)

    def test_repeated_cursor_fails(self):
        client = Mock()
        client.current_user_followed_artists.return_value = {
            "artists": {"items": [{"name": "a"}], "cursors": {"after": "same"}}
        }
        with self.assertRaises(RuntimeError):
            list(followed_artists(client))

    def test_genres_keep_cross_genre_membership(self):
        genres = group_by_genre([{"name": "One", "genres": ["house", "edm"]},
                                 {"name": "One", "genres": ["house"]}])
        self.assertEqual(genres, {"edm": ["One"], "house": ["One"]})
        self.assertEqual(genre_report(genres).count("  - One"), 2)
        self.assertEqual(extract_names(genres), ["One"])
        self.assertEqual(extract_names(genres, "HOUSE"), ["One"])

    def test_invalid_genre_data(self):
        for value in ([], {"items": [{"name": "One"}]}, {"house": "One"}):
            with self.assertRaises(ValueError):
                extract_names(value)

    def test_export_roundtrip_missing_fields_and_formula_text(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "artists.txt"
            export_artists([{"name": "=1+1", "genres": [],
                             "external_urls": {"spotify": "https://open.spotify.com/artist/test"}}], source)
            names, spreadsheet = convert_export(source)
            self.assertEqual(names.read_text(encoding="utf-8"), "=1+1\n")
            workbook = load_workbook(spreadsheet)
            sheet = workbook.active
            self.assertEqual(sheet["B2"].value, "=1+1")
            self.assertEqual(sheet["B2"].data_type, "s")
            self.assertIsNone(sheet["C2"].value)
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet["F2"].hyperlink.target, "https://open.spotify.com/artist/test")
            workbook.close()

    def test_empty_export(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "empty.txt"
            export_artists([], source)
            _, spreadsheet = convert_export(source)
            workbook = load_workbook(spreadsheet)
            self.assertEqual(workbook.active.max_row, 1)
            self.assertEqual(len(workbook.active.conditional_formatting), 0)
            workbook.close()

    def test_malformed_export_rejected(self):
        for content in ("", "Total artists: 1", "Total artists: 1\n\n1. Broken"):
            with self.assertRaises(ValueError):
                parse_export(content)

    def test_duplicate_recordings_require_artist_and_isrc(self):
        def track(identifier, artist="a", isrc="recording"):
            return {"id": identifier, "name": "Same title", "duration_ms": 100,
                    "artists": [{"id": artist}], "external_ids": {"isrc": isrc}}
        tracks = [track("1"), track("1"), track("2"), track("3", artist="other"),
                  track("4", isrc=None), track("5", isrc=None)]
        self.assertEqual([t["id"] for t in find_duplicates(tracks)], ["2"])

    def test_saved_track_pagination_and_missing_tracks(self):
        client = Mock()
        client.current_user_saved_tracks.side_effect = [
            {"items": [{"track": None}, {"track": {"id": "1"}}], "next": "next"},
            {"items": [{"track": {"id": "2", "is_local": True}}], "next": None},
        ]
        self.assertEqual(list(saved_tracks(client)), [{"id": "1"}])
        client.current_user_saved_tracks.assert_called_with(limit=50, offset=2)

    def test_playlist_matches_artist_and_deduplicates(self):
        client = Mock()
        client.search.side_effect = [
            {"artists": {"items": [{"name": "Artist", "id": "a"}]}},
            {"tracks": {"items": [
                {"uri": "spotify:track:1", "artists": [{"id": "a"}]},
                {"uri": "spotify:track:1", "artists": [{"id": "a"}]},
                {"uri": "spotify:track:2", "artists": [{"id": "b"}]},
            ]}},
        ]
        self.assertEqual(collect_tracks(client, ["Artist", "artist"]), ["spotify:track:1"])
        client.current_user_playlist_create.assert_not_called()

    def test_playlist_batches_and_empty_plan(self):
        client = Mock()
        with self.assertRaises(ValueError):
            create_playlist(client, "Empty", [])
        client.current_user_playlist_create.assert_not_called()
        client.current_user_playlist_create.return_value = {"id": "playlist"}
        create_playlist(client, "Test", [str(i) for i in range(205)])
        self.assertEqual([len(c.kwargs["payload"]["uris"]) for c in client._post.call_args_list],
                         [100, 100, 5])
        client.current_user_playlist_create.assert_called_once_with("Test", public=False)

    def test_scripts_import_without_authentication(self):
        root = Path(__file__).resolve().parents[1]
        with patch("spotipy.Spotify", side_effect=AssertionError("Unexpected authentication")):
            for path in root.glob("*.py"):
                spec = importlib.util.spec_from_file_location("check_" + path.stem, path)
                spec.loader.exec_module(importlib.util.module_from_spec(spec))

    def test_checked_in_genre_mapping(self):
        root = Path(__file__).resolve().parents[1]
        data = json.loads((root / "different_genres.json").read_text(encoding="utf-8"))
        self.assertTrue(extract_names(data))
        self.assertEqual((root / "genres-refined.txt").read_text(encoding="utf-8"),
                         genre_report(data))


if __name__ == "__main__":
    unittest.main()
